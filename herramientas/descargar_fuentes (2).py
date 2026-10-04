#!/usr/bin/env python3
"""
descargar_fuentes.py
====================

Descarga las páginas y documentos oficiales de los trámites estudiantiles de
PREGRADO de la Facultad de Ingeniería (UdeA), abarcando el Comité de Asuntos
Estudiantiles, Vicedecanatura (Supletorios y Segundo Calificador), Escuela de Idiomas
(PIFLE) y la normatividad del Reglamento Estudiantil. Extrae SOLO el contenido útil
(sin menús ni pies de página) y lo guarda con fecha, junto con un inventario en CSV.

Qué hace:
  1. Lee las URLs semilla de `semillas.txt`.
  2. Descarga cada página, respetando robots.txt y con pausa entre peticiones.
  3. Detecta el bloque de contenido principal y lo convierte a Markdown limpio.
  4. Sigue los enlaces que aparecen DENTRO de ese contenido (no los del menú),
     siempre que estén bajo los prefijos permitidos o sean documentos (PDF, Word).
  5. Extrae el texto de los PDF (y los campos si son formularios) y de los .docx.
  6. Guarda el original y el texto en fuentes/AAAA-MM-DD/ y actualiza
     fuentes/inventario.csv, marcando qué fuentes son nuevas o cambiaron.

Uso:
    pip install requests beautifulsoup4 pdfplumber pypdf python-docx
    python herramientas/descargar_fuentes.py
    python herramientas/descargar_fuentes.py --profundidad 2 --pausa 2
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import sys
import time
import unicodedata
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from urllib.parse import urldefrag, urljoin, urlparse
from urllib.robotparser import RobotFileParser

import requests
from bs4 import BeautifulSoup
from bs4.element import NavigableString, Tag

# --------------------------------------------------------------------------- #
# Configuración
# --------------------------------------------------------------------------- #

USER_AGENT = (
    "ChatbotUAE-ProyectoIntegrador-UdeA/1.0 "
    "(uso académico; recolección de información pública)"
)

# Prefijos autorizados para navegar enlaces HTML de las dependencias e instancias involucradas.
# Los documentos (PDF, Word) enlazados desde aquí se descargan aunque estén en /wps/wcm/connect/...
PREFIJOS_PERMITIDOS = [
    "https://www.udea.edu.co/wps/portal/udea/web/inicio/unidades-academicas/ingenieria/acerca-facultad/comites/comite-asuntos-estudiantiles",
    "https://www.udea.edu.co/wps/portal/udea/web/inicio/unidades-academicas/ingenieria/acerca-facultad/vicedecanatura",
    "https://www.udea.edu.co/wps/portal/udea/web/inicio/unidades-academicas/escuela-idiomas",
]

# Cualquier URL que contenga alguno de estos patrones se ignora automáticamente.
PATRONES_EXCLUIDOS = [
    "actas",            # Actas del Comité: contienen datos sensibles / casos particulares
    "login", "myportal", "registro-usuarios", "recuperar",
    "posgrado",         # El alcance del proyecto se limita exclusivamente a pregrado
    "javascript:", "mailto:", "tel:",
]

# Clases, IDs o etiquetas que representan ruido estructural (menús, pies de página, barras laterales).
RUIDO = re.compile(
    r"menu|nav|footer|cabezote|header|breadcrumb|signpost|social|compartir|"
    r"share|enlace|interes|cookie|banner|toolbar|accesib|idioma|search|buscar",
    re.IGNORECASE,
)

TIPOS_DOCUMENTO = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/msword": "doc",
}

# --------------------------------------------------------------------------- #
# Utilidades
# --------------------------------------------------------------------------- #


def normalizar_url(url: str) -> str:
    """Quita el fragmento (#...) y el estado dinámico del portal (/!ut/p/...),
    evitando duplicados en el inventario."""
    url, _ = urldefrag(url.strip())
    url = re.sub(r"/!ut/p/[^?]*", "", url)
    return url.rstrip("/")


def slug(texto: str, max_len: int = 60) -> str:
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    texto = re.sub(r"[^a-zA-Z0-9]+", "-", texto).strip("-").lower()
    return texto[:max_len].strip("-") or "sin-titulo"


def sha256(datos: bytes) -> str:
    return hashlib.sha256(datos).hexdigest()


def excluida(url: str) -> bool:
    u = url.lower()
    return any(p in u for p in PATRONES_EXCLUIDOS)


def permitida_html(url: str) -> bool:
    return any(url.startswith(p) for p in PREFIJOS_PERMITIDOS)


def parece_documento(url: str) -> bool:
    u = url.lower()
    return (
        u.endswith((".pdf", ".docx", ".doc"))
        or "/wps/wcm/connect/" in u  # Gestor de contenidos y documentos institucionales UdeA
    )


# --------------------------------------------------------------------------- #
# Extracción de contenido HTML
# --------------------------------------------------------------------------- #


def limpiar_ruido(soup: BeautifulSoup) -> None:
    for t in soup(["script", "style", "noscript", "form", "iframe", "svg",
                   "header", "nav", "footer", "button", "img"]):
        t.decompose()
    for t in list(soup.find_all(True)):
        if getattr(t, "decomposed", False) or t.attrs is None:
            continue
        marca = texto_atributo(t, "class") + " " + texto_atributo(t, "id")
        if RUIDO.search(marca) and densidad_enlaces(t) > 0.5:
            t.decompose()


def texto_atributo(t: Tag, nombre: str) -> str:
    valor = t.get(nombre)
    if valor is None:
        return ""
    if isinstance(valor, (list, tuple)):
        return " ".join(str(v) for v in valor)
    return str(valor)


def densidad_enlaces(t: Tag) -> float:
    total = len(t.get_text(" ", strip=True))
    if total == 0:
        return 1.0
    en_enlaces = sum(len(a.get_text(" ", strip=True)) for a in t.find_all("a"))
    return en_enlaces / total


def texto_sin_enlaces(t: Tag) -> int:
    total = len(t.get_text(" ", strip=True))
    en_enlaces = sum(len(a.get_text(" ", strip=True)) for a in t.find_all("a"))
    return total - 2 * en_enlaces


def bloque_principal(soup: BeautifulSoup) -> Tag:
    candidatos = soup.find_all(["main", "article", "section", "div"])
    if not candidatos:
        return soup.body or soup
    mejor = max(candidatos, key=texto_sin_enlaces)
    while True:
        hijos = [h for h in mejor.find_all(["div", "section", "article"], recursive=False)]
        if not hijos:
            break
        top = max(hijos, key=texto_sin_enlaces)
        if texto_sin_enlaces(top) >= 0.85 * texto_sin_enlaces(mejor):
            mejor = top
        else:
            break
    return mejor


def a_markdown(nodo: Tag, base: str) -> str:
    lineas: list[str] = []

    def en_linea(t) -> str:
        if isinstance(t, NavigableString):
            return str(t)
        if t.name == "a" and t.get("href"):
            txt = t.get_text(" ", strip=True)
            href = normalizar_url(urljoin(base, str(t["href"])))
            return f"[{txt}]({href})" if txt else ""
        if t.name in ("strong", "b"):
            txt = t.get_text(" ", strip=True)
            return f"**{txt}**" if txt else ""
        if t.name == "br":
            return "\n"
        return "".join(en_linea(h) for h in t.children)

    def bloque(t) -> None:
        if isinstance(t, NavigableString):
            s = str(t).strip()
            if s:
                lineas.append(s)
            return
        if not isinstance(t, Tag):
            return
        n = t.name
        if n in ("h1", "h2", "h3", "h4", "h5", "h6"):
            nivel = min(int(n[1]) + 1, 6)
            lineas.append("\n" + "#" * nivel + " " + t.get_text(" ", strip=True))
        elif n in ("p", "blockquote"):
            s = re.sub(r"[ \t]+", " ", en_linea(t)).strip()
            if s:
                lineas.append("\n" + s)
        elif n in ("ul", "ol"):
            for i, li in enumerate(t.find_all("li", recursive=False), 1):
                marca = f"{i}." if n == "ol" else "-"
                s = re.sub(r"\s+", " ", en_linea(li)).strip()
                if s:
                    lineas.append(f"{marca} {s}")
            lineas.append("")
        elif n == "table":
            filas = []
            for tr in t.find_all("tr"):
                celdas = [re.sub(r"\s+", " ", c.get_text(" ", strip=True))
                          for c in tr.find_all(["th", "td"])]
                if any(celdas):
                    filas.append("| " + " | ".join(celdas) + " |")
            if filas:
                ncol = filas[0].count("|") - 1
                filas.insert(1, "|" + " --- |" * ncol)
                lineas.append("\n" + "\n".join(filas) + "\n")
        else:
            for h in t.children:
                bloque(h)

    bloque(nodo)
    texto = "\n".join(lineas)
    return re.sub(r"\n{3,}", "\n\n", texto).strip()


def enlaces_de(nodo: Tag, base: str) -> list[str]:
    return [urljoin(base, str(a["href"])) for a in nodo.find_all("a", href=True)]


# --------------------------------------------------------------------------- #
# Extracción de documentos (PDF, Word)
# --------------------------------------------------------------------------- #


def texto_pdf(datos: bytes) -> tuple[str, str]:
    import pdfplumber
    from pypdf import PdfReader

    partes, obs = [], ""
    with pdfplumber.open(io.BytesIO(datos)) as pdf:
        for i, pagina in enumerate(pdf.pages, 1):
            t = (pagina.extract_text() or "").strip()
            if t:
                partes.append(f"<!-- página {i} -->\n{t}")
    texto = "\n\n".join(partes)
    if not texto:
        obs = "PDF escaneado o sin texto detectable: requiere verificación manual"

    try:
        campos = PdfReader(io.BytesIO(datos)).get_fields() or {}
        if campos:
            texto += "\n\n## Campos del formulario detectados\n" + "\n".join(
                f"- {nombre}" for nombre in campos)
    except Exception:
        pass
    return texto, obs


def texto_docx(datos: bytes) -> tuple[str, str]:
    import docx

    d = docx.Document(io.BytesIO(datos))
    partes = [p.text for p in d.paragraphs if p.text.strip()]
    for tabla in d.tables:
        for fila in tabla.rows:
            partes.append("| " + " | ".join(c.text.strip() for c in fila.cells) + " |")
    return "\n".join(partes), ""


# --------------------------------------------------------------------------- #
# Gestor de Descargas
# --------------------------------------------------------------------------- #


@dataclass
class Registro:
    url: str
    titulo: str = ""
    tipo: str = ""
    fecha_consulta: str = ""
    archivo_texto: str = ""
    archivo_original: str = ""
    sha256: str = ""
    palabras: int = 0
    cambio: str = ""
    estado: str = "ok"
    observaciones: str = ""


@dataclass
class Descargador:
    salida: Path
    pausa: float
    profundidad: int
    anteriores: dict[str, str] = field(default_factory=dict)
    visitadas: set[str] = field(default_factory=set)
    registros: list[Registro] = field(default_factory=list)
    robots: dict[str, RobotFileParser] = field(default_factory=dict)

    def __post_init__(self):
        self.hoy = date.today().isoformat()
        self.carpeta = self.salida / self.hoy
        self.carpeta.mkdir(parents=True, exist_ok=True)
        self.sesion = requests.Session()
        self.sesion.headers["User-Agent"] = USER_AGENT

    def puede(self, url: str) -> bool:
        p = urlparse(url)
        base = f"{p.scheme}://{p.netloc}"
        if base not in self.robots:
            rp = RobotFileParser()
            try:
                r = self.sesion.get(base + "/robots.txt", timeout=15)
                rp.parse(r.text.splitlines() if r.status_code == 200 else [])
            except requests.RequestException:
                rp.parse([])
            self.robots[base] = rp
        return self.robots[base].can_fetch(USER_AGENT, url)

    def obtener(self, url: str) -> requests.Response:
        ultimo: requests.RequestException = requests.RequestException(f"No se pudo obtener {url}")
        for intento in range(3):
            try:
                time.sleep(self.pausa)
                r = self.sesion.get(url, timeout=30)
                r.raise_for_status()
                return r
            except requests.RequestException as e:
                ultimo = e
                time.sleep(2 * (intento + 1))
        raise ultimo

    def guardar(self, reg: Registro, nombre: str, ext: str, original: bytes, texto: str) -> None:
        nombre = f"{slug(nombre)}-{sha256(reg.url.encode())[:6]}"
        (self.carpeta / f"{nombre}.{ext}").write_bytes(original)
        encabezado = (
            "---\n"
            f"titulo: \"{reg.titulo}\"\n"
            f"url: {reg.url}\n"
            f"tipo: {reg.tipo}\n"
            f"fecha_consulta: {reg.fecha_consulta}\n"
            f"sha256: {reg.sha256}\n"
            "---\n\n"
        )
        (self.carpeta / f"{nombre}.md").write_text(encabezado + texto, encoding="utf-8")
        reg.archivo_original = str((self.carpeta / f"{nombre}.{ext}").relative_to(self.salida))
        reg.archivo_texto = str((self.carpeta / f"{nombre}.md").relative_to(self.salida))
        reg.palabras = len(texto.split())
        previo = self.anteriores.get(reg.url)
        reg.cambio = "nuevo" if previo is None else (
            "sin cambios" if previo == reg.sha256 else "MODIFICADO")

    def procesar(self, url: str, nivel: int) -> None:
        url = normalizar_url(url)
        if url in self.visitadas or excluida(url):
            return
        self.visitadas.add(url)

        reg = Registro(url=url, fecha_consulta=self.hoy)
        self.registros.append(reg)
        print(f"[{nivel}] {url}")

        if not self.puede(url):
            reg.estado, reg.observaciones = "omitida", "bloqueada por robots.txt"
            return
        try:
            r = self.obtener(url)
        except requests.RequestException as e:
            reg.estado, reg.observaciones = "error", str(e)[:200]
            return

        ctype = r.headers.get("Content-Type", "").split(";")[0].strip().lower()
        reg.sha256 = sha256(r.content)

        try:
            if ctype in TIPOS_DOCUMENTO:
                reg.tipo = TIPOS_DOCUMENTO[ctype]
                texto, obs = (texto_pdf if reg.tipo == "pdf" else texto_docx)(r.content)
                reg.titulo = Path(urlparse(url).path).stem or "documento"
                reg.observaciones = obs
                self.guardar(reg, reg.titulo, reg.tipo, r.content, texto)

            elif ctype == "text/html":
                reg.tipo = "html"
                soup = BeautifulSoup(r.content, "html.parser")
                reg.titulo = (soup.title.get_text(strip=True) if soup.title else "") or url
                limpiar_ruido(soup)
                nodo = bloque_principal(soup)
                texto = a_markdown(nodo, url)
                reg.sha256 = sha256(texto.encode("utf-8"))
                if len(texto.split()) < 40:
                    reg.observaciones = "poco texto extraído: revisar manualmente"
                self.guardar(reg, reg.titulo, "html", r.content, texto)

                if nivel < self.profundidad:
                    for enlace in enlaces_de(nodo, url):
                        enlace = normalizar_url(enlace)
                        if permitida_html(enlace) or parece_documento(enlace):
                            self.procesar(enlace, nivel + 1)
            else:
                self.registros.remove(reg)
        except Exception as e:
            reg.estado, reg.observaciones = "error", f"al extraer: {e}"[:200]

    def escribir_inventario(self, ruta: Path) -> None:
        filas: dict[str, dict] = {}
        if ruta.exists():
            with ruta.open(encoding="utf-8") as f:
                for fila in csv.DictReader(f):
                    filas[fila["url"]] = fila
        for reg in self.registros:
            if reg.estado != "ok" and reg.url in filas:
                filas[reg.url]["estado"] = reg.estado
                filas[reg.url]["observaciones"] = reg.observaciones
            else:
                filas[reg.url] = reg.__dict__
        campos = list(Registro.__dataclass_fields__)
        with ruta.open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=campos)
            w.writeheader()
            for fila in sorted(filas.values(), key=lambda x: x["url"]):
                w.writerow({k: fila.get(k, "") for k in campos})


def leer_semillas(ruta: Path) -> list[str]:
    return [l.strip() for l in ruta.read_text(encoding="utf-8").splitlines()
            if l.strip() and not l.strip().startswith("#")]


def leer_hashes_previos(ruta: Path) -> dict[str, str]:
    if not ruta.exists():
        return {}
    with ruta.open(encoding="utf-8") as f:
        return {fila["url"]: fila["sha256"] for fila in csv.DictReader(f) if fila.get("sha256")}


def main() -> int:
    ap = argparse.ArgumentParser(description="Descarga fuentes oficiales para las fichas.")
    ap.add_argument("--semillas", type=Path, default=Path(__file__).with_name("semillas.txt"))
    ap.add_argument("--salida", type=Path, default=Path("fuentes"))
    ap.add_argument("--profundidad", type=int, default=1, help="Niveles de enlaces a seguir desde cada semilla (0 = solo semillas).")
    ap.add_argument("--pausa", type=float, default=1.5, help="Segundos de espera entre peticiones.")
    ap.add_argument("--prefijo", action="append", default=[], help="Prefijo de URL adicional permitido.")
    args = ap.parse_args()
    PREFIJOS_PERMITIDOS.extend(args.prefijo)

    args.salida.mkdir(parents=True, exist_ok=True)
    inventario = args.salida / "inventario.csv"

    d = Descargador(args.salida, args.pausa, args.profundidad, anteriores=leer_hashes_previos(inventario))
    for semilla in leer_semillas(args.semillas):
        d.procesar(semilla, 0)
    d.escribir_inventario(inventario)

    ok = sum(r.estado == "ok" for r in d.registros)
    cambios = [r for r in d.registros if r.cambio == "MODIFICADO"]
    revisar = [r for r in d.registros if r.observaciones]
    print(f"\nListo: {ok}/{len(d.registros)} fuentes guardadas en {d.carpeta}")
    print(f"Inventario: {inventario}")
    if cambios:
        print(f"\n{len(cambios)} fuente(s) CAMBIARON desde la última descarga:")
        for r in cambios:
            print("  -", r.url)
    if revisar:
        print(f"\n{len(revisar)} fuente(s) con observaciones:")
        for r in revisar:
            print(f"  - {r.url}\n    {r.observaciones}")
    return 0


if __name__ == "__main__":
    sys.exit(main())